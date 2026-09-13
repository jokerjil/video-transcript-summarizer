"""Build the downloadable deliverables from a finished pipeline result:

  * a single self-contained HTML document (summary + snapshots + transcript,
    images embedded as data URIs so the file stands alone),
  * a PDF (fpdf2, pure Python, no system dependency),
  * a Word document (python-docx, pure Python, no system dependency), and
  * a ZIP bundle of every raw artefact.

All of it runs offline - no network call, no external binary (no wkhtmltopdf,
no LibreOffice, no MS Word).
"""
from __future__ import annotations

import base64
import datetime as _dt
import html
import io
import math
import zipfile
from pathlib import Path

from util import iter_md_blocks, markdown_to_html, split_bold

_PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  body {{ font: 15px/1.6 system-ui, sans-serif; max-width: 860px; margin: 0 auto;
         padding: 32px 20px 80px; color: #1a1a1a; background: #fff; }}
  h1 {{ font-size: 24px; margin: 0 0 4px; }}
  h2 {{ font-size: 18px; margin: 28px 0 8px; border-bottom: 2px solid #e5e7eb; padding-bottom: 4px; }}
  .meta {{ color: #666; font-size: 13px; margin-bottom: 8px; }}
  ul {{ padding-left: 22px; }}
  li {{ margin: 4px 0; }}
  .frames {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(220px, 1fr));
            gap: 12px; margin-top: 10px; }}
  .frames figure {{ margin: 0; }}
  .frames img {{ width: 100%; border-radius: 6px; display: block; }}
  .frames figcaption {{ font-size: 12px; color: #555; margin-top: 3px; text-align: center; }}
  pre {{ white-space: pre-wrap; word-wrap: break-word; background: #f4f4f5; padding: 14px;
        border-radius: 8px; font: 12.5px/1.5 ui-monospace, monospace; }}
  @media print {{ .frames img {{ break-inside: avoid; }} h2 {{ break-after: avoid; }} }}
</style></head><body>
<h1>{title}</h1>
<div class="meta">{meta}</div>
{summary_html}
{frames_html}
<h2>Full Transcript</h2>
<pre>{transcript}</pre>
</body></html>
"""


def _meta_line(result: dict) -> str:
    bits = []
    src = result.get("source")
    if src:
        bits.append(f"Source: {html.escape(str(src))}")
    dur = result.get("duration")
    if dur:
        total = int(dur)
        bits.append(f"Duration: {total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}")
    if result.get("language"):
        bits.append(f"Language: {html.escape(str(result['language']))}")
    bits.append("Generated " + _dt.datetime.now().strftime("%Y-%m-%d %H:%M"))
    return " &nbsp;&middot;&nbsp; ".join(bits)


def _frames_html(snapshots: list) -> str:
    if not snapshots:
        return ""
    cells = []
    for shot in snapshots:
        p = Path(shot["file"])
        if not p.is_file():
            continue
        b64 = base64.b64encode(p.read_bytes()).decode("ascii")
        cells.append(
            f'<figure><img alt="frame at {shot["hms"]}" '
            f'src="data:image/jpeg;base64,{b64}">'
            f'<figcaption>{shot["hms"]}</figcaption></figure>'
        )
    if not cells:
        return ""
    return '<h2>Snapshots</h2>\n<div class="frames">\n' + "\n".join(cells) + "\n</div>"


def build_html(result: dict, title: str = "Video Summary") -> str:
    return _PAGE.format(
        title=html.escape(title),
        meta=_meta_line(result),
        summary_html=markdown_to_html(result.get("summary", "")),
        frames_html=_frames_html(result.get("snapshots") or []),
        transcript=html.escape(result.get("transcript") or "(no speech detected)"),
    )


def _meta_text(result: dict) -> str:
    bits = []
    src = result.get("source")
    if src:
        bits.append(f"Source: {src}")
    dur = result.get("duration")
    if dur:
        total = int(dur)
        bits.append(f"Duration: {total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}")
    if result.get("language"):
        bits.append(f"Language: {result['language']}")
    bits.append("Generated " + _dt.datetime.now().strftime("%Y-%m-%d %H:%M"))
    return "  |  ".join(bits)


def _pdf_safe(s: str) -> str:
    """fpdf2's built-in fonts only cover Latin-1; degrade anything else to '?'
    rather than crash the export. The HTML/ZIP downloads keep full Unicode."""
    return (s or "").encode("latin-1", "replace").decode("latin-1")


def build_pdf(result: dict, title: str = "Video Summary") -> bytes:
    from fpdf import FPDF

    pdf = FPDF(unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(18, 18, 18)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.multi_cell(0, 9, _pdf_safe(title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(110, 110, 110)
    pdf.multi_cell(0, 5, _pdf_safe(_meta_text(result)), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(3)

    def write_runs(text: str, size: float, indent: float = 0):
        pdf.set_x(pdf.l_margin + indent)
        for chunk, bold in split_bold(text):
            pdf.set_font("Helvetica", "B" if bold else "", size)
            pdf.write(5.2, _pdf_safe(chunk))
        pdf.ln(6.5)

    for kind, text in iter_md_blocks(result.get("summary", "")):
        if kind == "blank":
            pdf.ln(1.5)
        elif kind.startswith("h"):
            level = int(kind[1])
            size = {1: 15, 2: 13, 3: 11.5, 4: 10.5}.get(level, 10.5)
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", size)
            pdf.multi_cell(0, 7, _pdf_safe(text.replace("**", "")), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 10)
        elif kind == "li":
            write_runs("-  " + text, 10, indent=4)
        else:
            write_runs(text, 10)

    shots = [s for s in (result.get("snapshots") or []) if Path(s["file"]).is_file()]
    if shots:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Snapshots", new_x="LMARGIN", new_y="NEXT")
        cols, col_w, img_h, gap = 2, 82, 50, 9
        row_y = pdf.get_y()
        for i, shot in enumerate(shots):
            col = i % cols
            if col == 0:
                if row_y + img_h + 10 > 279:
                    pdf.add_page()
                    row_y = pdf.get_y()
            x = pdf.l_margin + col * (col_w + gap)
            pdf.image(shot["file"], x=x, y=row_y, w=col_w, h=img_h)
            pdf.set_xy(x, row_y + img_h + 1)
            pdf.set_font("Helvetica", "", 8)
            pdf.cell(col_w, 5, shot["hms"], align="C")
            if col == cols - 1 or i == len(shots) - 1:
                row_y += img_h + gap
        pdf.set_y(row_y)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Full Transcript", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9.5)
    pdf.multi_cell(0, 5, _pdf_safe(result.get("transcript") or "(no speech detected)"))

    return bytes(pdf.output())


def build_docx(result: dict, title: str = "Video Summary") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    doc.add_heading(title, level=0)
    meta_p = doc.add_paragraph(_meta_text(result))
    meta_p.runs[0].italic = True
    meta_p.runs[0].font.size = Pt(9)

    def add_runs(paragraph, text: str):
        for chunk, bold in split_bold(text):
            run = paragraph.add_run(chunk)
            run.bold = bold

    for kind, text in iter_md_blocks(result.get("summary", "")):
        if kind == "blank":
            continue
        if kind.startswith("h"):
            doc.add_heading(text.replace("**", ""), level=min(int(kind[1]), 4))
        elif kind == "li":
            add_runs(doc.add_paragraph(style="List Bullet"), text)
        else:
            add_runs(doc.add_paragraph(), text)

    shots = [s for s in (result.get("snapshots") or []) if Path(s["file"]).is_file()]
    if shots:
        doc.add_heading("Snapshots", level=1)
        cols = 2
        table = doc.add_table(rows=math.ceil(len(shots) / cols), cols=cols)
        for i, shot in enumerate(shots):
            cell = table.cell(i // cols, i % cols)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(shot["file"], width=Inches(2.7))
            cap = cell.add_paragraph(shot["hms"])
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size = Pt(8)

    doc.add_heading("Full Transcript", level=1)
    doc.add_paragraph(result.get("transcript") or "(no speech detected)")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_zip(result: dict, title: str = "video-summary") -> bytes:
    """Everything as a ZIP: every format plus the raw frame images."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("summary.md", result.get("summary", ""))
        z.writestr("transcript.txt", result.get("transcript") or "(no speech detected)")
        z.writestr("report.html", build_html(result, title))
        z.writestr("report.pdf", build_pdf(result, title))
        z.writestr("report.docx", build_docx(result, title))
        for shot in result.get("snapshots") or []:
            p = Path(shot["file"])
            if p.is_file():
                z.write(p, f"frames/{shot['name']}")
    return buf.getvalue()
